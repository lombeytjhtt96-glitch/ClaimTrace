import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette matching Healthkathon 2026
    DARK_NAVY = RGBColor(22, 38, 59)      # #16263B
    LIGHT_BG = RGBColor(248, 250, 252)    # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
    CARD_BORDER = RGBColor(167, 243, 208) # #A7F3D0 (Soft mint green)
    DARK_CARD_BG = RGBColor(15, 42, 58)   # #0F2A3A
    EMERALD = RGBColor(0, 176, 116)       # #00B074 (Theme green)
    EMERALD_DARK = RGBColor(16, 155, 99)  # #109B63
    MINT_LIGHT = RGBColor(236, 253, 245)  # #ECFDF5
    TEXT_DARK = RGBColor(30, 41, 59)      # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    TEXT_WHITE = RGBColor(255, 255, 255)
    ACCENT_CYAN = RGBColor(125, 168, 196) # #7DA8C4

    logo_path = '/home/Khalaf-Dev/.gemini/antigravity/scratch/claimtrace-jkn/healthkathon_logo.png'
    screenshot_path = '/home/Khalaf-Dev/.gemini/antigravity/brain/bc1217c1-62d5-40a9-8f47-eade18d108bf/screenshots/view_dual_cockpit_mobile.png'

    def add_header(slide, badge_text, title_text, category_text=None):
        # Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(0.85), Inches(0.85))
        badge.fill.solid()
        badge.fill.fore_color.rgb = EMERALD_DARK
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.size = Pt(26)
        p_b.font.bold = True
        p_b.font.color.rgb = TEXT_WHITE

        # Title
        title_box = slide.shapes.add_textbox(Inches(1.85), Inches(0.4), Inches(10.5), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_DARK

        # Subtitle / Category metadata if any
        if category_text:
            p_sub = tf_t.add_paragraph()
            p_sub.text = category_text
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_footer(slide, page_num):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.35))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Project Final Healthkathon 2026          Detect Smarter, Protect JKN - Smarter Protection for a Stronger JKN          Slide {page_num}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_NAVY
    bg1.line.fill.background()

    # Small top text
    top_box = s1.shapes.add_textbox(Inches(0.9), Inches(1.0), Inches(8), Inches(0.5))
    p_top = top_box.text_frame.paragraphs[0]
    p_top.text = "PANDUAN PROJECT FINAL"
    p_top.font.size = Pt(13)
    p_top.font.bold = True
    p_top.font.color.rgb = EMERALD

    # Main Title
    title_box1 = s1.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(9.5), Inches(1.2))
    p_main = title_box1.text_frame.paragraphs[0]
    p_main.text = "Healthkathon 2026"
    p_main.font.size = Pt(44)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_WHITE

    # Subtitle
    sub_box1 = s1.shapes.add_textbox(Inches(0.9), Inches(2.7), Inches(11), Inches(0.6))
    p_sub1 = sub_box1.text_frame.paragraphs[0]
    p_sub1.text = "Detect Smarter, Protect JKN - Smarter Protection for a Stronger JKN: Efisiensi Risiko Program JKN"
    p_sub1.font.size = Pt(13.5)
    p_sub1.font.italic = True
    p_sub1.font.color.rgb = ACCENT_CYAN

    # Logo
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(10.8), Inches(0.8), width=Inches(1.7))

    # 4 Info Boxes on Cover
    box_w = Inches(2.75)
    box_h = Inches(1.8)
    box_gap = Inches(0.24)
    start_x = Inches(0.9)
    start_y = Inches(4.3)

    cover_cards = [
        ("N A M A  S O L U S I", "ClaimTrace", "• Pre-Payment DSS\n• 3 Lapis Bukti Fraud"),
        ("N A M A  T I M", "TINEERENSE", "• MAN 1 Surakarta\n• Next-Gen E-Health Team"),
        ("I D  T I M", "HK2026-TINEERENSE", "• Kategori Faskes\n• Seleksi Finalist 2026"),
        ("K A T E G O R I", "Efisiensi Risiko Faskes", "• Sub #6 Phantom Billing\n• Sub #11 Repeat Billing")
    ]

    for i, (tag, val, desc) in enumerate(cover_cards):
        bx = start_x + i * (box_w + box_gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, start_y, box_w, box_h)
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD_BG
        card.line.color.rgb = EMERALD
        card.line.width = Pt(1.5)
        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.18)
        tf_c.margin_right = Inches(0.18)
        tf_c.margin_top = Inches(0.18)
        
        p1 = tf_c.paragraphs[0]
        p1.text = tag
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = EMERALD
        
        p2 = tf_c.add_paragraph()
        p2.text = val
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(4)
        
        p3 = tf_c.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = RGBColor(203, 213, 225)
        p3.space_before = Pt(4)

    # Note at bottom
    note_box = s1.shapes.add_textbox(Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.4))
    p_note = note_box.text_frame.paragraphs[0]
    p_note.text = "Inovasi Final Healthkathon 2026 — Sinergi Penapisan Pra-Bayar Multi-Lapis Berbasis Bukti (Human-in-the-Loop)"
    p_note.font.size = Pt(10)
    p_note.font.italic = True
    p_note.font.color.rgb = ACCENT_CYAN

    prs.save('/home/Khalaf-Dev/.gemini/antigravity/scratch/claimtrace-jkn/test_slide1.pptx')
    print('Slide 1 generated successfully!')

create_deck()
